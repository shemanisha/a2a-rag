from typing import List
from dataclasses import dataclass

from bs4 import BeautifulSoup
import csv

from ingestion.loader import LoadedPage, load_pdf
from docx import Document


@dataclass
class ExtractedDocument:
    pages: List[LoadedPage]
    metadata: dict


def extract_pdf(file_path: str) -> ExtractedDocument:
    pages = load_pdf(file_path)

    metadata = {"source": file_path, "format": "pdf"}

    return ExtractedDocument(pages=pages, metadata=metadata)


def extract_docx(file_path: str) -> ExtractedDocument:
    try:

        doc = Document(file_path)

        pages = []

        full_text = "\n".join(p.text for p in doc.paragraphs)

        pages.append(LoadedPage(text=full_text, page_number=0))

        metadata = {"source": file_path, "format": "docx"}

        return ExtractedDocument(pages=pages, metadata=metadata)

    except Exception:
        # As a safe fallback, raise so caller can handle
        raise


def extract_html(file_path: str) -> ExtractedDocument:
    with open(file_path, "r", encoding="utf-8") as fh:
        soup = BeautifulSoup(fh, "html.parser")

    text = soup.get_text(separator="\n")

    pages = [LoadedPage(text=text, page_number=0)]

    metadata = {"source": file_path, "format": "html"}

    return ExtractedDocument(pages=pages, metadata=metadata)


def extract_csv(file_path: str) -> ExtractedDocument:
    rows = []

    with open(file_path, newline="", encoding="utf-8") as csvfile:
        reader = csv.reader(csvfile)
        for row in reader:
            rows.append(" \t ".join(row))

    text = "\n".join(rows)

    pages = [LoadedPage(text=text, page_number=0)]

    metadata = {"source": file_path, "format": "csv"}

    return ExtractedDocument(pages=pages, metadata=metadata)


def extract(file_path: str, kind: str) -> ExtractedDocument:
    kind = kind.lower()

    if kind == "pdf":
        return extract_pdf(file_path)

    if kind == "docx":
        return extract_docx(file_path)

    if kind == "html":
        return extract_html(file_path)

    if kind == "csv":
        return extract_csv(file_path)

    raise ValueError(f"Unsupported extract kind: {kind}")
